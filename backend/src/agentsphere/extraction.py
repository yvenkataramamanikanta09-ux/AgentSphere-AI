"""Bounded, text-only document parsing, intended to run in the worker's child process.

These format checks reduce the parser's attack surface. They do not constitute a
malware scan. The worker must still enforce the wall-clock deadline by killing
the child process: a parser cannot reliably interrupt its own native-code work.
"""

from __future__ import annotations

import posixpath
import re
import stat
import time
import unicodedata
import zipfile
from contextlib import contextmanager
from io import BytesIO
from typing import Iterator
from urllib.parse import unquote, urlsplit

import pypdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from markdown_it import MarkdownIt
from pydantic import BaseModel, ConfigDict, Field, model_validator
from pypdf import PdfReader
from pypdf.errors import LimitReachedError
from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, StreamObject

PARSER_VERSION = "m5-stage3-v1"
_PDF_STREAM_BYTES = 8_000_000
_PDF_TOTAL_STREAM_BYTES = 40_000_000
_PDF_OBJECTS = 50_000
_STRUCTURE_NODES = 250_000
_STRUCTURE_DEPTH = 64
_TEXT_GROUP_LINES = 80
_TEXT_GROUP_CHARS = 16_000
_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_MAIN_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
_OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
_FORBIDDEN_PDF_KEYS = frozenset(
    {
        "/JS",
        "/JavaScript",
        "/AA",
        "/OpenAction",
        "/A",
        "/Launch",
        "/EmbeddedFiles",
        "/EF",
        "/AF",
        "/XFA",
        "/RichMedia",
        "/RichMediaContent",
        "/RichMediaSettings",
        "/Collection",
        "/SubmitForm",
        "/ImportData",
    }
)
_FORBIDDEN_PDF_ACTIONS = frozenset(
    {
        "/JavaScript",
        "/Launch",
        "/GoToR",
        "/GoToE",
        "/URI",
        "/SubmitForm",
        "/ImportData",
        "/Rendition",
        "/Movie",
        "/Sound",
        "/Hide",
        "/Named",
        "/ResetForm",
        "/SetOCGState",
    }
)


class ExtractionLimits(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    timeout_seconds: float = Field(default=60, gt=0, le=60)
    max_pdf_pages: int = Field(default=500, ge=1, le=500)
    max_text_bytes: int = Field(default=2_000_000, ge=1, le=2_000_000)
    max_sections: int = Field(default=5_000, ge=1, le=5_000)
    max_docx_entries: int = Field(default=1_024, ge=1, le=1_024)
    max_docx_expanded_bytes: int = Field(default=40_000_000, ge=1, le=40_000_000)
    max_docx_entry_bytes: int = Field(default=20_000_000, ge=1, le=20_000_000)
    max_docx_ratio: int = Field(default=100, ge=1, le=100)


class Section(BaseModel):
    model_config = ConfigDict(extra="forbid")

    text: str = Field(min_length=1)
    page_number: int | None = Field(default=None, ge=1, le=2_147_483_647)
    heading: str | None = Field(default=None, min_length=1, max_length=500)
    line_start: int | None = Field(default=None, ge=1, le=2_147_483_647)
    line_end: int | None = Field(default=None, ge=1, le=2_147_483_647)

    @model_validator(mode="after")
    def valid_source(self) -> Section:
        if not self.text.strip() or "\x00" in self.text:
            raise ValueError("Section text must be nonempty and contain no NUL")
        if self.heading is not None and (not self.heading.strip() or "\x00" in self.heading):
            raise ValueError("Section heading must be nonempty and contain no NUL")
        try:
            self.text.encode("utf-8", errors="strict")
            if self.heading is not None:
                self.heading.encode("utf-8", errors="strict")
        except UnicodeEncodeError as exc:
            raise ValueError("Section values must be valid Unicode") from exc
        if (self.line_start is None) != (self.line_end is None):
            raise ValueError("Line bounds must occur together")
        if self.line_start is not None and self.line_end < self.line_start:
            raise ValueError("Line bounds must be in source order")
        return self


class ExtractionResult(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sections: list[Section] = Field(max_length=5_000)
    warning_code: str | None = Field(default=None, pattern=r"^[a-z][a-z0-9_]{0,63}$")
    parser_version: str = Field(default=PARSER_VERSION, min_length=1, max_length=64)

    @model_validator(mode="after")
    def valid_codes(self) -> ExtractionResult:
        if self.warning_code is not None and not re.fullmatch(
            r"[a-z][a-z0-9_]{0,63}", self.warning_code
        ):
            raise ValueError("Invalid warning code")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}", self.parser_version):
            raise ValueError("Invalid parser version")
        if sum(len(section.text.encode("utf-8")) for section in self.sections) > 2_000_000:
            raise ValueError("Extracted text exceeds the publication limit")
        return self


class ExtractionError(Exception):
    """Only stable public codes cross the child-process boundary."""

    def __init__(self, code: str, retryable: bool = False) -> None:
        self.code = code
        self.retryable = retryable
        super().__init__(code)


class _Budget:
    def __init__(self, limits: ExtractionLimits) -> None:
        self.limits = limits
        self.deadline = time.monotonic() + limits.timeout_seconds
        self.sections: list[Section] = []
        self.text_bytes = 0

    def check(self) -> None:
        if time.monotonic() >= self.deadline:
            raise ExtractionError("extraction_timeout")

    def add(self, text: str, **source: object) -> None:
        self.check()
        # PostgreSQL text cannot store NUL. Other parser-produced controls have no
        # useful text meaning; plain text inputs are rejected instead of repaired.
        text = "".join(
            char for char in text if char in "\n\r\t" or unicodedata.category(char) != "Cc"
        ).strip()
        if not text:
            return
        heading = source.get("heading")
        if isinstance(heading, str):
            source["heading"] = heading.strip()[:500] or None
        self.text_bytes += len(text.encode("utf-8"))
        if self.text_bytes > self.limits.max_text_bytes:
            raise ExtractionError("text_limit_exceeded")
        if len(self.sections) >= self.limits.max_sections:
            raise ExtractionError("section_limit_exceeded")
        self.sections.append(Section(text=text, **source))

    def result(self) -> ExtractionResult:
        self.check()
        return ExtractionResult(
            sections=self.sections,
            warning_code=None if self.sections else "no_extractable_text",
        )


def extract_document(data: bytes, content_type: str, limits: ExtractionLimits) -> ExtractionResult:
    """Validate one of the four admitted formats and retain real source locations."""
    budget = _Budget(limits)
    media_type = content_type.partition(";")[0].strip().lower()
    handlers = {
        "application/pdf": _extract_pdf,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
        "text/plain": _extract_text,
        "text/markdown": _extract_markdown,
    }
    handler = handlers.get(media_type)
    if handler is None:
        raise ExtractionError("unsupported_type")
    try:
        handler(data, budget)
        return budget.result()
    except ExtractionError:
        raise
    except LimitReachedError as exc:
        raise ExtractionError("structure_limit_exceeded") from exc
    except Exception as exc:
        # Never put parser exception strings, document contents, or paths in jobs.
        code = "invalid_pdf" if media_type == "application/pdf" else "invalid_docx"
        if media_type.startswith("text/"):
            code = "invalid_text"
        raise ExtractionError(code) from exc


def _decode_text(data: bytes, budget: _Budget) -> str:
    if len(data) > budget.limits.max_text_bytes:
        raise ExtractionError("text_limit_exceeded")
    try:
        text = data.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise ExtractionError("invalid_utf8") from exc
    if any(unicodedata.category(char) == "Cc" and char not in "\t\n\r" for char in text):
        raise ExtractionError("unsafe_text")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _extract_text(data: bytes, budget: _Budget) -> None:
    text = _decode_text(data, budget)
    lines = text.split("\n")
    group: list[str] = []
    start = 1
    size = 0

    def flush(end: int) -> None:
        nonlocal group, size
        if group:
            budget.add("\n".join(group), line_start=start, line_end=end)
        group, size = [], 0

    for number, line in enumerate(lines, start=1):
        budget.check()
        if not line.strip():
            flush(number - 1)
            continue
        if group and (len(group) >= _TEXT_GROUP_LINES or size + len(line) > _TEXT_GROUP_CHARS):
            flush(number - 1)
        # Extremely long source lines remain traceable to the same real line.
        for offset in range(0, len(line), _TEXT_GROUP_CHARS):
            fragment = line[offset : offset + _TEXT_GROUP_CHARS]
            if not group:
                start = number
            group.append(fragment)
            size += len(fragment)
            if offset + _TEXT_GROUP_CHARS < len(line):
                flush(number)
    flush(len(lines))


def _inline_text(tokens: list, depth: int = 0) -> str:
    if depth > 32:
        raise ExtractionError("structure_limit_exceeded")
    pieces: list[str] = []
    for token in tokens:
        if token.type in {"text", "code_inline"}:
            pieces.append(token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            pieces.append("\n")
        elif token.type == "image":
            pieces.append(_inline_text(token.children or [], depth + 1) or token.content)
    return "".join(pieces)


def _extract_markdown(data: bytes, budget: _Budget) -> None:
    text = _decode_text(data, budget)
    # Parsing tokens never renders HTML or visits URLs. HTML-looking text is
    # ordinary text; image tokens contribute their alt text only.
    parser = MarkdownIt("commonmark", {"html": False, "maxNesting": 32})
    tokens = parser.parse(text)
    if len(tokens) > 100_000:
        raise ExtractionError("structure_limit_exceeded")
    heading: str | None = None
    in_heading = False
    for token in tokens:
        budget.check()
        if token.level + max(token.nesting, 0) >= 32:
            raise ExtractionError("structure_limit_exceeded")
        if token.type == "heading_open":
            in_heading = True
        elif token.type == "heading_close":
            in_heading = False
        elif token.type in {"inline", "fence", "code_block"} and token.map is not None:
            value = _inline_text(token.children or []) if token.type == "inline" else token.content
            if in_heading and value.strip():
                heading = value.strip()[:500]
            budget.add(value, heading=heading, line_start=token.map[0] + 1, line_end=token.map[1])


@contextmanager
def _pdf_configuration() -> Iterator[None]:
    # Limits must be installed BEFORE PdfReader: xref/object streams can already
    # decompress while the reader is opening the document.
    if hasattr(pypdf, "apply_configuration"):
        configuration = pypdf.get_configuration()
        candidates = {
            "maximum_declared_stream_length": _PDF_STREAM_BYTES,
            "array_based_stream_maximum_output_length": _PDF_STREAM_BYTES,
            "array_maximum_element_count": _PDF_OBJECTS,
            "dictionary_maximum_element_count": _PDF_OBJECTS,
            "zlib_maximum_output_length": _PDF_STREAM_BYTES,
            "lzw_maximum_output_length": _PDF_STREAM_BYTES,
            "run_length_maximum_output_length": _PDF_STREAM_BYTES,
            "brotli_maximum_output_length": _PDF_STREAM_BYTES,
            "stream_decoding_work_maximum_length": _PDF_TOTAL_STREAM_BYTES,
            "stream_parsing_work_maximum_count": _STRUCTURE_NODES,
            "page_tree_maximum_depth": _STRUCTURE_DEPTH,
            "page_tree_maximum_entries": 2_000,
            "xform_maximum_invocations_per_extraction": 256,
            "jbig2dec_binary": None,
        }
        supported = {key: value for key, value in candidates.items() if hasattr(configuration, key)}
        with pypdf.apply_configuration(**supported):
            yield
    else:
        # Older pypdf 6 releases expose these bounds as filter constants. Each
        # worker uses a separate process, and the previous values are restored.
        from pypdf import filters

        names = (
            "ZLIB_MAX_OUTPUT_LENGTH",
            "LZW_MAX_OUTPUT_LENGTH",
            "RUN_LENGTH_MAX_OUTPUT_LENGTH",
            "MAX_DECLARED_STREAM_LENGTH",
            "MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH",
        )
        previous = {name: getattr(filters, name) for name in names if hasattr(filters, name)}
        try:
            for name in previous:
                setattr(filters, name, _PDF_STREAM_BYTES)
            yield
        finally:
            for name, value in previous.items():
                setattr(filters, name, value)


def _extract_pdf(data: bytes, budget: _Budget) -> None:
    if not data.startswith(b"%PDF-") or not data.rstrip().endswith(b"%%EOF"):
        raise ExtractionError("invalid_pdf")
    with _pdf_configuration():
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ExtractionError("encrypted_pdf")
        references = [
            IndirectObject(number, generation, reader)
            for generation, entries in reader.xref.items()
            for number in entries
            if number != 0
        ]
        references.extend(IndirectObject(number, 0, reader) for number in reader.xref_objStm)
        if len(references) > _PDF_OBJECTS or int(reader.trailer.get("/Size", 0)) > _PDF_OBJECTS:
            raise ExtractionError("structure_limit_exceeded")
        active: set[tuple] = set()
        visited: set[tuple] = set()
        node_count = 0
        stream_bytes = 0
        page_count = 0

        def identity(value: object) -> tuple:
            if isinstance(value, IndirectObject):
                return ("reference", value.idnum, value.generation)
            return ("container", id(value))

        def visit(value: object, depth: int, *, glyph_map: bool = False) -> None:
            nonlocal node_count, stream_bytes, page_count
            budget.check()
            node_count += 1
            if depth > _STRUCTURE_DEPTH or node_count > _STRUCTURE_NODES:
                raise ExtractionError("structure_limit_exceeded")
            if not isinstance(value, (IndirectObject, DictionaryObject, ArrayObject)):
                if glyph_map:
                    raise ExtractionError("unsafe_pdf")
                return
            key = identity(value)
            if key in active:
                raise ExtractionError("unsafe_pdf")
            # A CharProcs dictionary has glyph names rather than action keys.
            # Revisit a shared dictionary when its semantic context changes.
            context_key = (*key, glyph_map)
            if context_key in visited:
                return
            active.add(key)
            if isinstance(value, IndirectObject):
                resolved = value.get_object()
                if resolved is None:
                    raise ExtractionError("invalid_pdf")
                visit(resolved, depth + 1, glyph_map=glyph_map)
            elif isinstance(value, DictionaryObject):
                if glyph_map:
                    if isinstance(value, StreamObject):
                        raise ExtractionError("unsafe_pdf")
                    for child in value.values():
                        resolved = (
                            child.get_object() if isinstance(child, IndirectObject) else child
                        )
                        if not isinstance(resolved, StreamObject):
                            raise ExtractionError("unsafe_pdf")
                        visit(child, depth + 1)
                    active.remove(key)
                    visited.add(context_key)
                    return
                kind = value.get("/Type")
                if kind in {"/Action", "/Filespec", "/EmbeddedFile"}:
                    raise ExtractionError("unsafe_pdf")
                if value.get("/S") in _FORBIDDEN_PDF_ACTIONS:
                    raise ExtractionError("unsafe_pdf")
                if kind == "/Pages" and int(value.get("/Count", 0)) > budget.limits.max_pdf_pages:
                    raise ExtractionError("pdf_page_limit")
                if kind == "/Page":
                    page_count += 1
                    if page_count > budget.limits.max_pdf_pages:
                        raise ExtractionError("pdf_page_limit")
                for name, child in value.items():
                    if name in _FORBIDDEN_PDF_KEYS:
                        raise ExtractionError("unsafe_pdf")
                    # A page tree has required parent links. Only a backlink to
                    # an ancestor is exempt; unrelated objects are still checked.
                    if name == "/Parent" and kind in {"/Page", "/Pages"}:
                        if identity(child) in active:
                            continue
                    if name == "/P" and kind == "/Annot" and identity(child) in active:
                        continue
                    if name == "/P" and kind == "/StructElem" and identity(child) in active:
                        parent = child.get_object() if isinstance(child, IndirectObject) else child
                        if (
                            parent is not value
                            and isinstance(parent, DictionaryObject)
                            and parent.get("/Type") in {"/StructTreeRoot", "/StructElem"}
                        ):
                            continue
                    visit(
                        child,
                        depth + 1,
                        glyph_map=(
                            name == "/CharProcs"
                            and kind == "/Font"
                            and value.get("/Subtype") == "/Type3"
                        ),
                    )
                if isinstance(value, StreamObject):
                    if len(value._data) > _PDF_STREAM_BYTES:
                        raise ExtractionError("structure_limit_exceeded")
                    # Image decoding is unnecessary for text extraction and can
                    # invoke optional native decoders; scanned pages stay empty.
                    if value.get("/Subtype") != "/Image":
                        filters = value.get("/Filter", [])
                        filters = filters if isinstance(filters, list) else [filters]
                        allowed = {
                            "/FlateDecode",
                            "/LZWDecode",
                            "/RunLengthDecode",
                            "/ASCII85Decode",
                            "/ASCIIHexDecode",
                        }
                        if len(filters) > 4 or any(item not in allowed for item in filters):
                            raise ExtractionError("unsafe_pdf")
                        decoded = value.get_data()
                        stream_bytes += len(decoded)
                        if (
                            len(decoded) > _PDF_STREAM_BYTES
                            or stream_bytes > _PDF_TOTAL_STREAM_BYTES
                        ):
                            raise ExtractionError("structure_limit_exceeded")
            else:
                if glyph_map:
                    raise ExtractionError("unsafe_pdf")
                if len(value) > _PDF_OBJECTS:
                    raise ExtractionError("structure_limit_exceeded")
                for child in value:
                    visit(child, depth + 1)
            active.remove(key)
            visited.add(context_key)

        visit(reader.trailer, 0)
        for reference in references:
            # The xref sweep finds unreachable objects. Objects already checked
            # through their owner retain that owner's semantic context here.
            if not any((*identity(reference), context) in visited for context in (False, True)):
                visit(reference, 0)
        if len(reader.pages) > budget.limits.max_pdf_pages:
            raise ExtractionError("pdf_page_limit")
        for number, page in enumerate(reader.pages, start=1):
            budget.check()
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > _PDF_STREAM_BYTES:
                raise ExtractionError("structure_limit_exceeded")
            budget.add(page.extract_text() or "", page_number=number)


def _safe_xml(data: bytes, budget: _Budget) -> etree._Element:
    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        dtd_validation=False,
        huge_tree=False,
        recover=False,
    )
    root = etree.fromstring(data, parser=parser)
    if root.getroottree().docinfo.doctype:
        raise ExtractionError("unsafe_docx")
    count = 0
    stack = [(root, 0)]
    while stack:
        node, depth = stack.pop()
        budget.check()
        count += 1
        if count > _STRUCTURE_NODES or depth > _STRUCTURE_DEPTH:
            raise ExtractionError("structure_limit_exceeded")
        if isinstance(node, etree._Entity):
            raise ExtractionError("unsafe_docx")
        if isinstance(node.tag, str) and etree.QName(node).localname in {
            "altChunk",
            "OLEObject",
            "object",
            "attachedTemplate",
        }:
            raise ExtractionError("unsafe_docx")
        stack.extend((child, depth + 1) for child in node)
    return root


def _relationship_target(part: str, target: str) -> str:
    value = unquote(target)
    parsed = urlsplit(value)
    if parsed.scheme or parsed.netloc or "\\" in value or "\x00" in value or parsed.query:
        raise ExtractionError("unsafe_docx")
    base = "" if part == "_rels/.rels" else posixpath.dirname(posixpath.dirname(part))
    path = posixpath.normpath(posixpath.join(base, parsed.path))
    if path.startswith("/"):
        path = path[1:]
    if path == ".." or path.startswith("../") or path in {"", "."}:
        raise ExtractionError("unsafe_docx")
    return path


def _docx_preflight(data: bytes, budget: _Budget) -> None:
    xml_parts: dict[str, etree._Element] = {}
    members: dict[str, bytes] = {}
    with zipfile.ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > budget.limits.max_docx_entries:
            raise ExtractionError("docx_expansion_limit")
        names: set[str] = set()
        folded: set[str] = set()
        total = 0
        for entry in entries:
            budget.check()
            name = entry.filename
            path = name.rstrip("/")
            if (
                not path
                or "\\" in name
                or name.startswith("/")
                or ":" in name
                or any(piece in {"", ".", ".."} for piece in path.split("/"))
                or "\x00" in name
                or entry.orig_filename != name
                or name.casefold() in folded
                or stat.S_ISLNK(entry.external_attr >> 16)
                or entry.flag_bits & 1
                or entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
            ):
                raise ExtractionError("unsafe_docx")
            folded.add(name.casefold())
            names.add(name)
            lower = name.casefold()
            if any(
                marker in lower
                for marker in (
                    "vbaproject",
                    "vbadata",
                    "embeddings/",
                    "activex/",
                    "oleobject",
                )
            ):
                raise ExtractionError("unsafe_docx")
            if (
                entry.file_size > budget.limits.max_docx_entry_bytes
                or entry.file_size > max(entry.compress_size, 1) * budget.limits.max_docx_ratio
                or total + entry.file_size > budget.limits.max_docx_expanded_bytes
            ):
                raise ExtractionError("docx_expansion_limit")
            chunks: list[bytes] = []
            actual = 0
            with archive.open(entry) as member:
                while chunk := member.read(65_536):
                    budget.check()
                    actual += len(chunk)
                    total += len(chunk)
                    if (
                        actual > budget.limits.max_docx_entry_bytes
                        or total > budget.limits.max_docx_expanded_bytes
                        or actual > max(entry.compress_size, 1) * budget.limits.max_docx_ratio
                    ):
                        raise ExtractionError("docx_expansion_limit")
                    chunks.append(chunk)
            if actual != entry.file_size:
                raise ExtractionError("invalid_docx")
            members[name] = b"".join(chunks)
            if members[name].startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                raise ExtractionError("unsafe_docx")
            if lower.endswith((".xml", ".rels")):
                xml_parts[name] = _safe_xml(members[name], budget)

    required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
    if not required.issubset(xml_parts):
        raise ExtractionError("invalid_docx")
    types = xml_parts["[Content_Types].xml"]
    if types.tag != f"{{{_CONTENT_NS}}}Types":
        raise ExtractionError("invalid_docx")
    main_found = False
    xml_extensions: set[str] = set()
    xml_overrides: set[str] = set()
    for declaration in types:
        content = declaration.get("ContentType", "")
        if any(
            marker in content.casefold() for marker in ("macroenabled", "vbaproject", "oleobject")
        ):
            raise ExtractionError("unsafe_docx")
        if declaration.get("PartName") == "/word/document.xml" and content == _MAIN_TYPE:
            main_found = True
        if content.endswith("+xml") or content in {"application/xml", "text/xml"}:
            if declaration.get("PartName"):
                xml_overrides.add(declaration.get("PartName").lstrip("/").lower())
            elif declaration.get("Extension"):
                xml_extensions.add(declaration.get("Extension").casefold())
    # PartFactory dispatches on content type, not filename extension. Inspect an
    # XML part even when an attacker gives it a misleading non-XML suffix.
    for name, member in members.items():
        if name not in xml_parts and (
            name.lower() in xml_overrides or name.rsplit(".", 1)[-1].casefold() in xml_extensions
        ):
            xml_parts[name] = _safe_xml(member, budget)
    if not main_found:
        raise ExtractionError("invalid_docx")
    office_targets = []
    for part, root in xml_parts.items():
        if not part.endswith(".rels"):
            continue
        if root.tag != f"{{{_REL_NS}}}Relationships":
            raise ExtractionError("invalid_docx")
        ids: set[str] = set()
        for relationship in root:
            if relationship.tag != f"{{{_REL_NS}}}Relationship":
                raise ExtractionError("invalid_docx")
            rel_id = relationship.get("Id", "")
            rel_type = relationship.get("Type", "")
            if not rel_id or rel_id in ids:
                raise ExtractionError("invalid_docx")
            ids.add(rel_id)
            if relationship.get("TargetMode", "Internal").casefold() != "internal":
                raise ExtractionError("unsafe_docx")
            relation_kind = rel_type.rsplit("/", 1)[-1].casefold()
            if relation_kind in {
                "afchunk",
                "oleobject",
                "package",
                "vbaproject",
                "attachedtemplate",
                "control",
            }:
                raise ExtractionError("unsafe_docx")
            target = _relationship_target(part, relationship.get("Target", ""))
            if target not in names:
                raise ExtractionError("invalid_docx")
            if part == "_rels/.rels" and rel_type == _OFFICE_REL:
                office_targets.append(target)
    main = xml_parts["word/document.xml"]
    if (
        office_targets != ["word/document.xml"]
        or main.tag != f"{{{_WORD_NS}}}document"
        or main.find(f"{{{_WORD_NS}}}body") is None
    ):
        raise ExtractionError("invalid_docx")


def _word_cell_text(cell: etree._Element) -> str:
    parts: list[str] = []
    for element in cell.iter():
        if element.tag == f"{{{_WORD_NS}}}t":
            parts.append(element.text or "")
        elif element.tag == f"{{{_WORD_NS}}}tab":
            parts.append("\t")
        elif element.tag in {f"{{{_WORD_NS}}}br", f"{{{_WORD_NS}}}cr"}:
            parts.append("\n")
        elif element.tag == f"{{{_WORD_NS}}}p" and parts:
            parts.append("\n")
    return "".join(parts).strip()


def _extract_docx(data: bytes, budget: _Budget) -> None:
    _docx_preflight(data, budget)
    document = Document(BytesIO(data))
    paragraph_number = table_number = 0
    heading: str | None = None

    def reference(label: str) -> str:
        return (label + (f" · {heading}" if heading else ""))[:500]

    # python-docx's separate .paragraphs and .tables collections lose their
    # interleaving; iter_inner_content retains the actual body order.
    for block in document.iter_inner_content():
        budget.check()
        if isinstance(block, Paragraph):
            paragraph_number += 1
            value = block.text
            style = block.style.name.casefold() if block.style is not None else ""
            if value.strip() and (re.fullmatch(r"heading\s*\d+", style) or style == "title"):
                heading = value.strip()[:500]
            budget.add(value, heading=reference(f"Paragraph {paragraph_number}"))
        elif isinstance(block, Table):
            table_number += 1
            for row_number, row in enumerate(block._tbl.tr_lst, start=1):
                # Read physical cells once so merged cells are not duplicated.
                value = "\t".join(_word_cell_text(cell) for cell in row.tc_lst)
                budget.add(value, heading=reference(f"Table {table_number}, row {row_number}"))
