"""Live HTTP + separate worker verification in a newly created disposable database.

Never connects to/migrates the application database. Generated accounts, originals,
logs, and database belong only to this run and are removed in finally.
"""

import json
import os
import secrets
import socket
import subprocess
import sys
import sysconfig
import time
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import UUID, uuid4

import httpx
import psycopg
from docx import Document as WordDocument
from psycopg import sql
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import create_engine, select
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session

from agentsphere.config import BACKEND_DIR, Settings
from agentsphere.models import DocumentFile, DocumentJob, DocumentSection, JobKind, JobState
from agentsphere.storage import PrivateStorage, new_storage_key


def module_command(module, *arguments):
    if os.name == "nt":
        bootstrap = (
            "import runpy,site,sys;"
            f"site.addsitedir({sysconfig.get_paths()['purelib']!r});"
            f"sys.argv={[module, *arguments]!r};"
            f"runpy.run_module({module!r},run_name='__main__')"
        )
        return [sys._base_executable, "-I", "-c", bootstrap]
    return [sys.executable, "-m", module, *arguments]


def pdf_bytes(encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 24 100 Td (Verified PDF evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt(secrets.token_urlsafe(16))
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def word_bytes():
    document = WordDocument()
    document.add_heading("Verified research", 1)
    document.add_paragraph("A real DOCX paragraph.")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "Source", "Evidence"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def main():
    settings = Settings()
    if settings.test_database_url is None or settings.database_url is None:
        raise SystemExit("Existing separate local development/test URL configuration is required.")
    template = make_url(settings.test_database_url.get_secret_value())
    development = make_url(settings.database_url.get_secret_value())
    for url in (template, development):
        if url.host not in {"localhost", "127.0.0.1", "::1"} or url.query:
            raise SystemExit("Only query-free loopback database URLs are supported.")
    if not (template.database or "").endswith("_test") or template.database == development.database:
        raise SystemExit("A distinct local test database template is required.")
    cluster = json.loads((BACKEND_DIR / ".local/postgres-config.json").read_text())
    if cluster["port"] != template.port:
        raise SystemExit("Local cluster/test port mismatch; nothing changed.")
    name = "agentsphere_processing_" + uuid4().hex + "_test"
    url = template.set(database=name).render_as_string(hide_password=False)
    admin_options = dict(
        host=template.host,
        port=template.port,
        dbname="postgres",
        user="agentsphere_admin",
        password=cluster["admin_password"],
        autocommit=True,
        connect_timeout=3,
    )
    created, server, engine = False, None, None
    passed = 0

    def check(label, condition):
        nonlocal passed
        if not condition:
            raise AssertionError(label)
        passed += 1
        print("PASS " + label, flush=True)

    try:
        with psycopg.connect(**admin_options) as admin:
            admin.execute(
                sql.SQL("CREATE DATABASE {} OWNER {} TEMPLATE template0").format(
                    sql.Identifier(name), sql.Identifier(template.username)
                )
            )
            created = True
            admin.execute(
                sql.SQL("REVOKE CONNECT ON DATABASE {} FROM PUBLIC").format(sql.Identifier(name))
            )
        print("Created a new disposable processing database; development is untouched.", flush=True)
        with TemporaryDirectory(prefix="agentsphere_processing_") as temporary:
            root = Path(temporary)
            with socket.socket() as probe:
                probe.bind(("127.0.0.1", 0))
                port = probe.getsockname()[1]
            origin = f"http://127.0.0.1:{port}"
            environment = os.environ | {
                "DATABASE_URL": url,
                "TEST_DATABASE_URL": url,
                "PRIVATE_STORAGE_ROOT": str(root / "private"),
                "APP_ORIGIN": origin,
                "ALLOW_INSECURE_LOCAL_HTTP": "true",
                "ENVIRONMENT": "development",
                "AUTH_RATE_LIMIT_KEY": secrets.token_urlsafe(48),
            }
            with (root / "process.log").open("w", encoding="utf-8") as log:
                for arguments in (("upgrade", "head"), ("check",)):
                    subprocess.run(
                        module_command("alembic", *arguments),
                        cwd=BACKEND_DIR,
                        env=environment,
                        check=True,
                        stdout=log,
                        stderr=log,
                    )
                server = subprocess.Popen(
                    module_command(
                        "uvicorn",
                        "agentsphere.main:create_app",
                        "--factory",
                        "--host",
                        "127.0.0.1",
                        "--port",
                        str(port),
                        "--no-proxy-headers",
                    ),
                    cwd=BACKEND_DIR,
                    env=environment,
                    stdout=log,
                    stderr=log,
                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
                )
                deadline = time.monotonic() + 30
                while True:
                    if server.poll() is not None:
                        raise RuntimeError("Disposable API failed to start.")
                    try:
                        if httpx.get(origin + "/health/live", timeout=1).status_code == 200:
                            break
                    except httpx.TransportError:
                        pass
                    if time.monotonic() >= deadline:
                        raise RuntimeError("Disposable API startup timeout.")
                    time.sleep(0.1)
                check("real FastAPI health", True)
                engine = create_engine(url, hide_parameters=True)
                clients = []

                def worker(*arguments):
                    return subprocess.run(
                        module_command("agentsphere.worker", *arguments),
                        cwd=BACKEND_DIR,
                        env=environment,
                        stdout=log,
                        stderr=log,
                        timeout=90,
                        check=False,
                    ).returncode

                def upload(client, base, filename, payload, mime, nonce=None):
                    nonce = nonce or uuid4()
                    response = client.post(
                        base + "/uploads",
                        files={"file": (filename, payload, mime)},
                        headers={"Idempotency-Key": str(nonce)},
                    )
                    check("upload accepted " + filename, response.status_code == 202)
                    return response.json(), nonce

                try:
                    for role in ("owner", "editor", "viewer", "outsider"):
                        client = httpx.Client(
                            base_url=origin, headers={"Origin": origin}, timeout=15
                        )
                        clients.append(client)
                        bootstrap = client.get("/api/v1/auth/csrf")
                        check(role + " CSRF bootstrap", bootstrap.status_code == 200)
                        client.headers["X-CSRF-Token"] = bootstrap.json()["csrf_token"]
                        identity = {
                            "email": f"processing-{role}-{uuid4().hex}@example.com",
                            "password": secrets.token_urlsafe(32),
                            "display_name": role,
                        }
                        registered = client.post("/api/v1/auth/register", json=identity)
                        check(role + " registration", registered.status_code == 201)
                        client.user_id = registered.json()["user"]["id"]
                        login = client.post(
                            "/api/v1/auth/login",
                            json={"email": identity["email"], "password": identity["password"]},
                        )
                        check(role + " login and session rotation", login.status_code == 200)
                        client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
                    owner, editor, viewer, outsider = clients
                    workspace = owner.post(
                        "/api/v1/workspaces", json={"name": "Worker verification"}
                    )
                    check("workspace creation", workspace.status_code == 201)
                    workspace_id = UUID(workspace.json()["id"])
                    base = f"/api/v1/workspaces/{workspace_id}/documents"
                    for member, role in ((editor, "editor"), (viewer, "viewer")):
                        response = owner.post(
                            f"/api/v1/workspaces/{workspace_id}/members",
                            json={"user_id": member.user_id, "role": role},
                        )
                        check(role + " membership", response.status_code == 201)
                    formats = [
                        ("evidence.pdf", pdf_bytes(), "application/pdf"),
                        (
                            "evidence.docx",
                            word_bytes(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        ),
                        (
                            "evidence.txt",
                            b"First source line.\nSecond source line.\n",
                            "text/plain",
                        ),
                        (
                            "evidence.md",
                            b"# Evidence\n\nA real Markdown paragraph.\n",
                            "text/markdown",
                        ),
                    ]
                    records = []
                    for filename, payload, mime in formats:
                        body, nonce = upload(owner, base, filename, payload, mime)
                        detail = base + "/" + body["document"]["id"]
                        check(
                            "queued downloads/sections blocked " + filename,
                            owner.get(detail + "/download").status_code == 409
                            and owner.get(detail + "/sections").status_code == 409,
                        )
                        check("separate worker process " + filename, worker("--once") == 0)
                        metadata = owner.get(detail + "/file").json()
                        check(
                            "ready with successful processing " + filename,
                            metadata["processing_status"] == "ready"
                            and metadata["download_available"],
                        )
                        sections = owner.get(detail + "/sections").json()["items"]
                        check(
                            "published source references " + filename,
                            bool(sections)
                            and (
                                sections[0]["page_number"] == 1
                                if filename.endswith(".pdf")
                                else sections[0]["heading"] is not None
                                if filename.endswith(".docx")
                                else sections[0]["line_start"] == 1
                            ),
                        )
                        check(
                            "verified original download " + filename,
                            owner.get(detail + "/download").content == payload,
                        )
                        records.append((body, nonce, payload, mime))
                    detail = base + "/" + records[0][0]["document"]["id"]
                    for member, role in ((owner, "owner"), (editor, "editor"), (viewer, "viewer")):
                        check(
                            role + " authorized sections",
                            member.get(detail + "/sections").status_code == 200,
                        )
                    check(
                        "outsider sections rejected",
                        outsider.get(detail + "/sections").status_code == 404,
                    )
                    check(
                        "viewer cannot delete processed document",
                        viewer.delete(detail).status_code == 403,
                    )
                    replay, _ = upload(owner, base, *formats[0], nonce=records[0][1])
                    check(
                        "lost-response retry reuses processed record",
                        replay["replayed"]
                        and replay["document"]["id"] == records[0][0]["document"]["id"],
                    )
                    check("empty worker poll succeeds", worker("--once") == 0)
                    for filename, payload in (
                        ("broken.pdf", b"%PDF-1.7\nbroken\n%%EOF"),
                        ("encrypted.pdf", pdf_bytes(encrypted=True)),
                    ):
                        body, _ = upload(owner, base, filename, payload, "application/pdf")
                        detail = base + "/" + body["document"]["id"]
                        check("invalid original worker completes safely", worker("--once") == 0)
                        check(
                            "invalid/encrypted original failed without publication",
                            owner.get(detail + "/file").json()["processing_status"] == "failed"
                            and owner.get(detail + "/sections").status_code == 409
                            and owner.get(detail + "/download").status_code == 409,
                        )
                    body, _ = upload(
                        owner, base, "retry.txt", b"Transient storage recovery.\n", "text/plain"
                    )
                    with Session(engine) as db:
                        attachment = db.get(DocumentFile, UUID(body["file"]["id"]))
                        key = attachment.storage_key
                    original = root / "private" / key
                    held = root / "held-original"
                    original.rename(held)
                    check("transient storage failure handled", worker("--once") == 0)
                    check(
                        "transient error queued for bounded retry",
                        owner.get(base + "/" + body["document"]["id"] + "/file").json()[
                            "processing_status"
                        ]
                        == "queued",
                    )
                    held.rename(original)
                    with Session(engine) as db, db.begin():
                        job = db.scalar(select(DocumentJob).where(DocumentJob.storage_key == key))
                        job.available_at = datetime.now(UTC) - timedelta(seconds=1)
                    check("separate worker retries restored original", worker("--once") == 0)
                    check(
                        "retry published successfully",
                        owner.get(base + "/" + body["document"]["id"] + "/file").json()[
                            "processing_status"
                        ]
                        == "ready",
                    )
                    body, _ = upload(
                        owner, base, "crash.txt", b"Durable crash recovery.\n", "text/plain"
                    )
                    crash = (
                        "import os;from sqlalchemy.orm import sessionmaker;"
                        "from agentsphere.config import Settings;"
                        "from agentsphere.database import build_engine;"
                        "from agentsphere.worker_jobs import DurableJobQueue;"
                        "q=DurableJobQueue(sessionmaker(build_engine(Settings().database_url.get_secret_value())));"
                        "q.claim();os._exit(77)"
                    )
                    bootstrap = (
                        "import site;" + f"site.addsitedir({sysconfig.get_paths()['purelib']!r});"
                    )
                    executable = sys._base_executable if os.name == "nt" else sys.executable
                    crashed = subprocess.run(
                        [executable, "-I", "-c", bootstrap + crash],
                        env=environment,
                        cwd=BACKEND_DIR,
                        stdout=log,
                        stderr=log,
                        timeout=15,
                        check=False,
                    )
                    check("worker process crash after durable claim", crashed.returncode == 77)
                    with Session(engine) as db, db.begin():
                        job = db.scalar(
                            select(DocumentJob).where(
                                DocumentJob.document_id == UUID(body["document"]["id"])
                            )
                        )
                        check("crash leaves durable running lease", job.state == JobState.RUNNING)
                        job.started_at = datetime.now(UTC) - timedelta(seconds=180)
                        job.lease_expires_at = datetime.now(UTC) - timedelta(seconds=1)
                    check("replacement worker recovers stale lease", worker("--once") == 0)
                    check(
                        "recovered original ready",
                        owner.get(base + "/" + body["document"]["id"] + "/file").json()[
                            "processing_status"
                        ]
                        == "ready",
                    )
                    body, nonce, payload, mime = records[2]
                    with Session(engine) as db:
                        key = db.get(DocumentFile, UUID(body["file"]["id"])).storage_key
                    detail = base + "/" + body["document"]["id"]
                    check(
                        "editor deletes processed metadata",
                        editor.delete(detail).status_code == 204,
                    )
                    check(
                        "deletion immediately hides APIs",
                        owner.get(detail).status_code == 404
                        and owner.get(detail + "/sections").status_code == 404,
                    )
                    check(
                        "physical original retained until cleanup",
                        (root / "private" / key).is_file(),
                    )
                    check("separate cleanup worker", worker("--once") == 0)
                    check(
                        "cleanup physically removed original", not (root / "private" / key).exists()
                    )
                    rejected = owner.post(
                        base + "/uploads",
                        files={"file": ("evidence.txt", payload, mime)},
                        headers={"Idempotency-Key": str(nonce)},
                    )
                    check(
                        "cleanup preserves deleted-upload nonce tombstone",
                        rejected.status_code == 409,
                    )
                    storage = PrivateStorage(root / "private")
                    storage.initialize()
                    orphan = storage.put(
                        new_storage_key(workspace_id), BytesIO(b"Old orphan"), max_bytes=100
                    )
                    orphan_path = root / "private" / orphan.key
                    old = time.time() - 172_800
                    os.utime(orphan_path, (old, old))
                    scratch = root / "private/tmp" / (uuid4().hex + ".part")
                    scratch.write_bytes(b"Stale scratch")
                    os.utime(scratch, (old, old))
                    with storage.activity():
                        check("active upload fence blocks maintenance", worker("--reconcile") == 1)
                        check("active fence retains old temporary", scratch.is_file())
                    check("separate reconciliation pass", worker("--reconcile") == 0)
                    check("stale temporary removed", not scratch.exists())
                    check("orphan cleanup queued before unlink", orphan_path.exists())
                    check("orphan cleanup worker", worker("--once") == 0)
                    check("orphan physically removed", not orphan_path.exists())
                    with Session(engine) as db:
                        check(
                            "terminal cleanup tombstones retained",
                            bool(
                                db.scalars(
                                    select(DocumentJob).where(
                                        DocumentJob.kind == JobKind.DELETE_FILES,
                                        DocumentJob.state == JobState.SUCCEEDED,
                                    )
                                ).all()
                            ),
                        )
                        check(
                            "extracted sections persist in PostgreSQL",
                            bool(db.scalars(select(DocumentSection)).all()),
                        )
                finally:
                    for client in clients:
                        client.close()
                    if server is not None:
                        server.terminate()
                        server.wait(timeout=10)
                        server = None
                    if engine is not None:
                        engine.dispose()
                        engine = None
        print(f"{passed} real API/PostgreSQL/worker checks passed; 0 failed.", flush=True)
        return 0
    finally:
        if server is not None:
            server.terminate()
            server.wait(timeout=10)
        if engine is not None:
            engine.dispose()
        if created:
            with psycopg.connect(**admin_options) as admin:
                admin.execute(sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(name)))
            print("Removed only this run's disposable processing database.", flush=True)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (psycopg.Error, subprocess.CalledProcessError):
        raise SystemExit(
            "Disposable processing verification failed; no development migration attempted."
        ) from None
