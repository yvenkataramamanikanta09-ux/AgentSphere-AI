"""API uploads processed by the real worker, parser subprocess, and PostgreSQL."""

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Event
from uuid import UUID, uuid4

import pytest
from conftest import register_and_login
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import event, func, select, text, update
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from agentsphere.extraction import ExtractionResult, Section, extract_document
from agentsphere.models import (
    Document,
    DocumentFile,
    DocumentJob,
    DocumentSection,
    JobKind,
    JobState,
    ProcessingStatus,
)
from agentsphere.worker import ProcessingWorker
from agentsphere.worker_jobs import DurableJobQueue

pytestmark = pytest.mark.postgres
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def processing_client(client_factory, tmp_path):
    return register_and_login(client_factory(private_storage_root=tmp_path / "private"))


def upload(
    client,
    payload=b"Research notes.\n",
    filename="notes.txt",
    content_type="text/plain",
    *,
    nonce=None,
    url=None,
):
    response = client.post(
        (url or client.documents_url) + "/uploads",
        files={"file": (filename, payload, content_type)},
        headers={"Idempotency-Key": str(nonce or uuid4())},
    )
    assert response.status_code == 202, response.text
    return response.json()


def worker(client, *, extractor=None, stopping=None):
    return ProcessingWorker(
        DurableJobQueue(client.app.state.session_factory),
        client.app.state.private_storage,
        client.app.state.settings,
        extractor=extractor,
        stopping=stopping,
    )


def path(client, body):
    return f"{client.documents_url}/{body['document']['id']}"


def assert_unavailable(client, location, *, status):
    attachment = client.get(location + "/file")
    assert attachment.status_code == 200, attachment.text
    assert attachment.json()["processing_status"] == status
    assert attachment.json()["download_available"] is False
    assert client.get(location + "/sections").status_code == 409
    assert client.get(location + "/download").status_code == 409
    return attachment.json()


def make_available(client, document_id):
    with client.app.state.session_factory() as db, db.begin():
        db.execute(
            update(DocumentJob)
            .where(DocumentJob.document_id == UUID(document_id))
            .values(available_at=func.clock_timestamp() - text("interval '1 second'"))
        )


def pdf_bytes(*pages, password=None):
    writer = PdfWriter()
    for value in pages:
        page = writer.add_blank_page(width=300, height=300)
        if value is not None:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 12 Tf 10 100 Td ({value}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(content)
    if password is not None:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes():
    document = WordDocument()
    document.add_heading("Research", 1)
    document.add_paragraph("First paragraph")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Table evidence"
    row.cells[1].text = "Confirmed"
    document.add_paragraph("Last paragraph")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("kind", ["txt", "md", "pdf", "docx"])
def test_actual_upload_worker_sections_download_and_replay(processing_client, db_env, kind):
    client = processing_client
    fixtures = {
        "txt": (b"First line\nSecond line\n\nFinal line\n", "text/plain"),
        "md": (b"# Research\n\nEvidence **confirmed**.\n", "text/markdown"),
        "pdf": (pdf_bytes("Page one", None, "Page three"), PDF),
        "docx": (docx_bytes(), DOCX),
    }
    payload, content_type = fixtures[kind]
    nonce = uuid4()
    body = upload(client, payload, f"research.{kind}", content_type, nonce=nonce)
    location = path(client, body)
    assert_unavailable(client, location, status="queued")
    with Session(db_env) as db:
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
    processing = worker(client)
    assert processing.extractor is None  # This path invokes the killable parser subprocess.
    assert processing.run_once()
    attachment = client.get(location + "/file").json()
    assert attachment["processing_status"] == "ready" and attachment["download_available"]
    assert attachment["processed_at"] is not None and attachment["error_code"] is None
    assert "storage_key" not in attachment and "parser_version" not in attachment
    response = client.get(location + "/sections")
    assert response.status_code == 200, response.text
    page = response.json()
    items = page["items"]
    assert page["document_id"] == body["document"]["id"] and not page["has_more"]
    assert [item["ordinal"] for item in items] == list(range(len(items)))
    if kind == "txt":
        assert [(item["text"], item["line_start"], item["line_end"]) for item in items] == [
            ("First line\nSecond line", 1, 2),
            ("Final line", 4, 4),
        ]
    elif kind == "md":
        assert [item["text"] for item in items] == ["Research", "Evidence confirmed."]
        assert [(item["line_start"], item["line_end"]) for item in items] == [(1, 1), (3, 3)]
        assert all(item["heading"] == "Research" for item in items)
    elif kind == "pdf":
        assert [(item["text"], item["page_number"]) for item in items] == [
            ("Page one", 1),
            ("Page three", 3),
        ]
        assert all(item["line_start"] is None and item["line_end"] is None for item in items)
    else:
        assert [item["text"] for item in items] == [
            "Research",
            "First paragraph",
            "Table evidence\tConfirmed",
            "Last paragraph",
        ]
        assert [item["heading"] for item in items] == [
            "Paragraph 1 · Research",
            "Paragraph 2 · Research",
            "Table 1, row 1 · Research",
            "Paragraph 3 · Research",
        ]
        assert all(item["page_number"] is None for item in items)
    if kind != "pdf":
        assert all(item["page_number"] is None for item in items)
    download = client.get(location + "/download")
    assert download.status_code == 200 and download.content == payload
    assert download.headers["content-type"].startswith(content_type)
    assert download.headers["content-length"] == str(len(payload))
    assert download.headers["content-disposition"].startswith("attachment;")
    assert download.headers["cache-control"] == "no-store"
    assert download.headers["x-content-type-options"] == "nosniff"
    replay = upload(client, payload, f"research.{kind}", content_type, nonce=nonce)
    assert replay["replayed"] and replay["document"]["id"] == body["document"]["id"]
    assert replay["file"]["processing_status"] == "ready"
    assert client.get(location + "/sections").json()["items"] == items
    assert not processing.run_once()
    with Session(db_env) as db:
        jobs = db.scalars(select(DocumentJob)).all()
        assert len(jobs) == 1 and jobs[0].state == JobState.SUCCEEDED
        assert jobs[0].attempts == 1


@pytest.mark.parametrize("role", ["editor", "viewer"])
def test_processed_content_respects_membership_and_revocation(
    processing_client, client_factory, role
):
    owner = processing_client
    payload = b"Private workspace evidence.\n"
    body = upload(owner, payload)
    assert worker(owner).run_once()
    member = register_and_login(
        client_factory(private_storage_root=owner.app.state.settings.private_storage_root)
    )
    members_url = f"/api/v1/workspaces/{owner.workspace_id}/members"
    assert (
        owner.post(members_url, json={"user_id": member.user_id, "role": role}).status_code == 201
    )
    location = path(owner, body)
    assert member.get(location + "/file").json()["download_available"]
    assert member.get(location + "/sections").json()["items"][0]["text"] == payload.decode().strip()
    assert member.get(location + "/download").content == payload
    if role == "viewer":
        response = member.post(
            owner.documents_url + "/uploads",
            files={"file": ("forbidden.txt", b"Viewer write", "text/plain")},
            headers={"Idempotency-Key": str(uuid4())},
        )
        assert response.status_code == 403
        assert member.delete(location).status_code == 403
    assert owner.delete(f"{members_url}/{member.user_id}").status_code == 204
    for suffix in ("/file", "/sections", "/download"):
        assert member.get(location + suffix).status_code == 404


def test_processed_content_cannot_cross_workspaces_or_authentication(
    processing_client, client_factory
):
    owner = processing_client
    body = upload(owner)
    assert worker(owner).run_once()
    outsider = register_and_login(client_factory())
    anonymous = client_factory()
    owner_location = path(owner, body)
    wrong_workspace = f"{outsider.documents_url}/{body['document']['id']}"
    for suffix in ("/file", "/sections", "/download"):
        assert outsider.get(owner_location + suffix).status_code == 404
        assert outsider.get(wrong_workspace + suffix).status_code == 404
        assert anonymous.get(owner_location + suffix).status_code == 401
        assert owner.get(owner_location + suffix).status_code == 200


@pytest.mark.parametrize("failure", ["malformed", "encrypted"])
def test_pdf_processing_failure_is_terminal_and_never_exposes_sections(
    processing_client, db_env, failure
):
    client = processing_client
    payload = (
        b"%PDF-1.7\n1 0 obj\n<<>>\nendobj\n%%EOF\n"
        if failure == "malformed"
        else pdf_bytes("Encrypted evidence", password="secret")
    )
    body = upload(client, payload, "research.pdf", PDF)
    location = path(client, body)
    assert_unavailable(client, location, status="queued")
    assert worker(client).run_once()
    attachment = assert_unavailable(client, location, status="failed")
    assert attachment["error_code"] == (
        "invalid_pdf" if failure == "malformed" else "encrypted_pdf"
    )
    assert "secret" not in str(attachment)
    with Session(db_env) as db:
        job = db.scalars(select(DocumentJob)).one()
        assert job.state == JobState.FAILED and job.attempts == 1
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
        assert db.get(Document, UUID(body["document"]["id"])).status.value == "registered"


def test_blank_pdf_is_ready_with_explicit_warning_and_original_available(processing_client, db_env):
    client = processing_client
    payload = pdf_bytes(None)
    body = upload(client, payload, "scanned.pdf", PDF)
    assert worker(client).run_once()
    location = path(client, body)
    attachment = client.get(location + "/file").json()
    assert attachment["processing_status"] == "ready" and attachment["download_available"]
    assert attachment["warning_code"] == "no_extractable_text"
    page = client.get(location + "/sections").json()
    assert page["items"] == [] and not page["has_more"]
    assert client.get(location + "/download").content == payload
    with Session(db_env) as db:
        assert db.scalars(select(DocumentJob.state)).one() == JobState.SUCCEEDED


def test_sections_paginate_real_text_groups_without_inventing_pages(processing_client):
    client = processing_client
    lines = [f"Evidence line {number}" for number in range(1, 242)]
    body = upload(client, "\n".join(lines).encode())
    assert worker(client).run_once()
    location = path(client, body) + "/sections"
    first = client.get(location, params={"limit": 2, "offset": 0}).json()
    second = client.get(location, params={"limit": 2, "offset": 2}).json()
    assert first["has_more"] and not second["has_more"]
    items = first["items"] + second["items"]
    assert [item["ordinal"] for item in items] == [0, 1, 2, 3]
    assert [(item["line_start"], item["line_end"]) for item in items] == [
        (1, 80),
        (81, 160),
        (161, 240),
        (241, 241),
    ]
    assert "\n".join(item["text"] for item in items).splitlines() == lines
    assert all(item["page_number"] is None for item in items)
    empty = client.get(location, params={"limit": 2, "offset": 4}).json()
    assert empty["items"] == [] and not empty["has_more"]
    for params in ({"limit": 0}, {"limit": 101}, {"offset": -1}):
        assert client.get(location, params=params).status_code == 422


def test_api_deletion_during_parsing_cannot_resurrect_document_or_results(
    processing_client, db_env
):
    client = processing_client
    body = upload(client)
    location = path(client, body)
    parsing, resume = Event(), Event()

    def slow_parser(data, content_type):
        parsing.set()
        assert resume.wait(5)
        return extract_document(data, content_type, client.app.state.settings.extraction_limits)

    processing = worker(client, extractor=slow_parser)
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(processing.run_once)
        try:
            assert parsing.wait(3)
            assert_unavailable(client, location, status="processing")
            assert client.delete(location).status_code == 204
        finally:
            resume.set()
        assert future.result(timeout=5)
    for suffix in ("", "/file", "/sections", "/download"):
        assert client.get(location + suffix).status_code == 404
    assert processing.run_once()  # The independent cleanup job removes the original.
    assert not processing.run_once()
    with Session(db_env) as db:
        assert db.scalar(select(func.count()).select_from(Document)) == 0
        assert db.scalar(select(func.count()).select_from(DocumentFile)) == 0
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
        history = db.scalars(select(DocumentJob)).all()
        assert len(history) == 2 and all(job.document_id is None for job in history)
        assert {job.state for job in history} == {JobState.CANCELLED, JobState.SUCCEEDED}
    assert list(client.app.state.private_storage.root.glob("objects/*/*/original")) == []


def test_worker_stop_before_parsing_requeues_then_new_worker_recovers(processing_client, db_env):
    client = processing_client
    body = upload(client)
    stopping = Event()
    stopping.set()

    def forbidden_parser(*args):
        pytest.fail("A stopped worker invoked a parser")

    assert worker(client, extractor=forbidden_parser, stopping=stopping).run_once()
    location = path(client, body)
    assert_unavailable(client, location, status="queued")
    with Session(db_env) as db:
        job = db.scalars(select(DocumentJob)).one()
        assert job.state == JobState.QUEUED and job.attempts == 1
        assert job.error_code == "worker_interrupted"
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
    make_available(client, body["document"]["id"])
    assert worker(client).run_once()
    assert client.get(location + "/file").json()["processing_status"] == "ready"


def test_dead_lease_drops_late_result_and_another_worker_recovers(processing_client, db_env):
    client = processing_client
    body = upload(client, b"Current original.\n")
    document_id = UUID(body["document"]["id"])

    def stale_parser(*args):
        with client.app.state.session_factory() as db, db.begin():
            db.execute(
                update(DocumentJob)
                .where(DocumentJob.document_id == document_id)
                .values(
                    started_at=func.clock_timestamp() - text("interval '5 minutes'"),
                    lease_expires_at=func.clock_timestamp() - text("interval '1 second'"),
                )
            )
        return ExtractionResult(sections=[Section(text="Late stale result")])

    assert worker(client, extractor=stale_parser).run_once()
    location = path(client, body)
    assert_unavailable(client, location, status="processing")
    with Session(db_env) as db:
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
    assert worker(client).run_once()
    assert client.get(location + "/sections").json()["items"][0]["text"] == "Current original."
    with Session(db_env) as db:
        job = db.scalars(select(DocumentJob)).one()
        assert job.state == JobState.SUCCEEDED and job.attempts == 2


def test_database_publication_failure_has_no_partial_api_results_and_retries(
    processing_client, db_env
):
    client = processing_client
    body = upload(client, b"First result\n\nSecond result\n")

    def reject_publication(connection, cursor, statement, parameters, context, executemany):
        if statement.lstrip().lower().startswith("insert into document_sections"):
            raise OperationalError("redacted", {}, Exception("Injected database detail"))

    event.listen(db_env, "after_cursor_execute", reject_publication)
    try:
        assert worker(client).run_once()
    finally:
        event.remove(db_env, "after_cursor_execute", reject_publication)
    location = path(client, body)
    assert_unavailable(client, location, status="queued")
    with Session(db_env) as db:
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
        job = db.scalars(select(DocumentJob)).one()
        assert job.state == JobState.QUEUED and job.error_code == "database_unavailable"
    make_available(client, body["document"]["id"])
    assert worker(client).run_once()
    items = client.get(location + "/sections").json()["items"]
    assert [item["text"] for item in items] == ["First result", "Second result"]


@pytest.mark.parametrize("damage", ["same_size", "oversized"])
def test_original_integrity_mismatch_fails_before_parser_and_publishes_nothing(
    processing_client, db_env, damage
):
    client = processing_client
    payload = b"Private original evidence.\n"
    body = upload(client, payload)
    with Session(db_env) as db:
        key = db.get(DocumentFile, UUID(body["file"]["id"])).storage_key
    original = client.app.state.private_storage.object_path(key)
    original.write_bytes(b"X" * len(payload) if damage == "same_size" else payload + b"X")

    def forbidden_parser(*args):
        pytest.fail("An original with invalid integrity reached the parser")

    assert worker(client, extractor=forbidden_parser).run_once()
    attachment = assert_unavailable(client, path(client, body), status="failed")
    assert attachment["error_code"] == "original_integrity_failed"
    with Session(db_env) as db:
        assert db.scalar(select(func.count()).select_from(DocumentSection)) == 0
        assert db.scalars(select(DocumentJob.state)).one() == JobState.FAILED


def test_processing_does_not_change_archived_metadata_or_renamed_filename(
    processing_client, db_env
):
    client = processing_client
    body = upload(client)
    location = path(client, body)
    response = client.patch(location, json={"status": "archived", "filename": "renamed.txt"})
    assert response.status_code == 200, response.text
    archived = response.json()
    assert worker(client).run_once()
    assert client.get(location).json() == archived
    assert client.get(location + "/file").json()["processing_status"] == "ready"
    download = client.get(location + "/download")
    assert download.status_code == 200 and "renamed.txt" in download.headers["content-disposition"]
    with Session(db_env) as db:
        assert db.scalars(select(DocumentJob.kind)).one() == JobKind.EXTRACT
        assert db.scalars(select(DocumentFile.processing_status)).one() == ProcessingStatus.READY
