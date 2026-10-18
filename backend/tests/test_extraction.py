"""Synthetic format fixtures exercise parsing without documents or network access."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from docx import Document
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf._page import PageObject
from pypdf.generic import (
    ArrayObject,
    BooleanObject,
    DecodedStreamObject,
    DictionaryObject,
    FloatObject,
    NameObject,
    NumberObject,
    TextStringObject,
)

from agentsphere import extraction
from agentsphere.extraction import (
    ExtractionError,
    ExtractionLimits,
    ExtractionResult,
    Section,
    extract_document,
)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
LIMITS = ExtractionLimits()


def pdf_fixture(*pages: str | None, change=None, password: str | None = None) -> bytes:
    writer = PdfWriter()
    for value in pages:
        page = writer.add_blank_page(width=300, height=300)
        if value is not None:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 10 100 Td ({value}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    if change:
        change(writer)
    if password is not None:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_fixture() -> bytes:
    document = Document()
    document.add_heading("Plan", 1)
    document.add_paragraph("First paragraph")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "First table"
    document.add_paragraph("Second paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Second table"
    table.cell(0, 1).text = "Beside it"
    table.cell(1, 0).text = "Last row"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def type3_pdf_fixture(change=None, *, indirect=True) -> bytes:
    def add_type3_font(writer):
        font = writer.pages[0]["/Resources"]["/Font"]["/F1"]
        glyph = DecodedStreamObject()
        glyph.set_data(b"600 0 0 0 600 700 d1 0 0 600 700 re f")
        glyphs = DictionaryObject({NameObject("/A"): writer._add_object(glyph)})
        if indirect:
            glyphs = writer._add_object(glyphs)
        font.update(
            {
                NameObject("/Subtype"): NameObject("/Type3"),
                NameObject("/FontBBox"): ArrayObject([NumberObject(n) for n in (0, 0, 600, 700)]),
                NameObject("/FontMatrix"): ArrayObject(
                    [
                        FloatObject(0.001),
                        NumberObject(0),
                        NumberObject(0),
                        FloatObject(0.001),
                        NumberObject(0),
                        NumberObject(0),
                    ]
                ),
                NameObject("/CharProcs"): glyphs,
                NameObject("/Encoding"): DictionaryObject(
                    {NameObject("/Differences"): ArrayObject([NumberObject(65), NameObject("/A")])}
                ),
                NameObject("/FirstChar"): NumberObject(65),
                NameObject("/LastChar"): NumberObject(65),
                NameObject("/Widths"): ArrayObject([NumberObject(600)]),
                NameObject("/Resources"): DictionaryObject(),
            }
        )
        if change:
            change(writer, glyphs, glyph)

    return pdf_fixture("A", change=add_type3_font)


def change_zip(data: bytes, replacements: dict[str, bytes], extras=()) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name not in replacements or replacements[name] is not None:
                target.writestr(name, replacements.get(name, source.read(name)))
        for name, value in extras:
            target.writestr(name, value)
    return output.getvalue()


def assert_error(data: bytes, content_type: str, code: str, **limits) -> ExtractionError:
    with pytest.raises(ExtractionError) as caught:
        extract_document(data, content_type, ExtractionLimits(**limits))
    assert caught.value.code == code
    return caught.value


def test_pdf_extracts_real_page_numbers_and_skips_blank_pages():
    result = extract_document(pdf_fixture("One", None, "Three"), PDF, LIMITS)
    assert [(section.text, section.page_number) for section in result.sections] == [
        ("One", 1),
        ("Three", 3),
    ]
    assert all(section.line_start is None for section in result.sections)
    assert result.parser_version == "m5-stage3-v1"


def test_blank_pdf_is_ready_with_warning_and_no_ocr():
    result = extract_document(pdf_fixture(None), PDF, LIMITS)
    assert result.sections == []
    assert result.warning_code == "no_extractable_text"


@pytest.mark.parametrize("password", ["secret", ""])
def test_encrypted_pdf_is_rejected_even_with_empty_password(password):
    assert_error(pdf_fixture(None, password=password), PDF, "encrypted_pdf")


@pytest.mark.parametrize("data", [b"hello.pdf", b"%PDF-1.7\nnot a PDF\n%%EOF", b"%PDF-1.7"])
def test_invalid_pdf_is_not_accepted_as_plain_text(data):
    assert_error(data, PDF, "invalid_pdf")


@pytest.mark.parametrize(
    "key",
    [
        "/JS",
        "/JavaScript",
        "/AA",
        "/OpenAction",
        "/A",
        "/Launch",
        "/EmbeddedFiles",
        "/EF",
        "/AF",
        "/XFA",
        "/RichMedia",
        "/Collection",
    ],
)
def test_pdf_rejects_active_features_in_dictionary_keys(key):
    data = pdf_fixture(
        "Text",
        change=lambda writer: writer._root_object.update(
            {NameObject(key): TextStringObject("unused")}
        ),
    )
    assert_error(data, PDF, "unsafe_pdf")


def test_pdf_also_inspects_unreachable_action_objects():
    def add_action(writer):
        writer._add_object(DictionaryObject({NameObject("/S"): NameObject("/JavaScript")}))

    assert_error(pdf_fixture("Text", change=add_action), PDF, "unsafe_pdf")


def test_pdf_rejects_cycles_while_allowing_normal_page_parent_links():
    def cycle(writer):
        node = DictionaryObject()
        reference = writer._add_object(node)
        node[NameObject("/Self")] = reference
        writer._root_object[NameObject("/Custom")] = reference

    assert_error(pdf_fixture("Text", change=cycle), PDF, "unsafe_pdf")
    assert extract_document(pdf_fixture("Text"), PDF, LIMITS).sections[0].text == "Text"


@pytest.mark.parametrize("indirect", [False, True])
def test_pdf_type3_glyph_names_are_not_mistaken_for_actions(indirect):
    result = extract_document(type3_pdf_fixture(indirect=indirect), PDF, LIMITS)
    assert result.sections[0].text == "A"


@pytest.mark.parametrize("shared_first", [False, True])
def test_pdf_type3_glyph_context_does_not_exempt_shared_action_keys(shared_first):
    def share(writer, glyphs, glyph):
        original = list(writer._root_object.items())
        if shared_first:
            writer._root_object.clear()
        writer._root_object[NameObject("/Custom")] = glyphs
        if shared_first:
            writer._root_object.update(original)

    assert_error(type3_pdf_fixture(share), PDF, "unsafe_pdf")


def test_pdf_type3_glyph_streams_still_reject_action_dictionaries():
    def add_action(writer, glyphs, glyph):
        glyph[NameObject("/A")] = DictionaryObject({NameObject("/S"): NameObject("/JavaScript")})

    assert_error(type3_pdf_fixture(add_action), PDF, "unsafe_pdf")


def test_pdf_type3_glyph_map_rejects_non_stream_action_values():
    def add_action(writer, glyphs, glyph):
        glyphs.get_object()[NameObject("/A")] = DictionaryObject(
            {NameObject("/S"): NameObject("/JavaScript")}
        )

    assert_error(type3_pdf_fixture(add_action), PDF, "unsafe_pdf")


def test_pdf_accepts_tagged_structure_parent_backlinks():
    def add_structure(writer):
        root = DictionaryObject({NameObject("/Type"): NameObject("/StructTreeRoot")})
        root_ref = writer._add_object(root)
        group = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/StructElem"),
                NameObject("/S"): NameObject("/Document"),
                NameObject("/P"): root_ref,
            }
        )
        group_ref = writer._add_object(group)
        paragraph_ref = writer._add_object(
            DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/StructElem"),
                    NameObject("/S"): NameObject("/P"),
                    NameObject("/P"): group_ref,
                    NameObject("/Pg"): writer.pages[0].indirect_reference,
                    NameObject("/K"): NumberObject(0),
                }
            )
        )
        group[NameObject("/K")] = paragraph_ref
        root[NameObject("/K")] = group_ref
        root[NameObject("/ParentTree")] = writer._add_object(
            DictionaryObject(
                {NameObject("/Nums"): ArrayObject([NumberObject(0), ArrayObject([paragraph_ref])])}
            )
        )
        writer._root_object[NameObject("/StructTreeRoot")] = root_ref
        writer._root_object[NameObject("/MarkInfo")] = DictionaryObject(
            {NameObject("/Marked"): BooleanObject(True)}
        )
        page = writer.pages[0]
        page[NameObject("/StructParents")] = NumberObject(0)
        stream = page["/Contents"]
        stream.set_data(b"/P <</MCID 0>> BDC " + stream.get_data() + b" EMC")

    result = extract_document(pdf_fixture("Accessible text", change=add_structure), PDF, LIMITS)
    assert result.sections[0].text == "Accessible text"


def test_pdf_structural_parent_exemption_does_not_allow_unrelated_cycles():
    def add_cycle(writer):
        node = DictionaryObject({NameObject("/Type"): NameObject("/StructElem")})
        reference = writer._add_object(node)
        node[NameObject("/P")] = writer._root_object.indirect_reference
        writer._root_object[NameObject("/Custom")] = reference

    assert_error(pdf_fixture("Text", change=add_cycle), PDF, "unsafe_pdf")


def test_pdf_structural_parent_exemption_does_not_allow_self_cycles():
    def add_cycle(writer):
        node = DictionaryObject({NameObject("/Type"): NameObject("/StructElem")})
        reference = writer._add_object(node)
        node[NameObject("/P")] = reference
        writer._root_object[NameObject("/Custom")] = reference

    assert_error(pdf_fixture("Text", change=add_cycle), PDF, "unsafe_pdf")


def test_pdf_page_count_is_checked_before_text_extraction(monkeypatch):
    data = pdf_fixture("One", "Two")
    monkeypatch.setattr(PageObject, "extract_text", lambda *args, **kwargs: pytest.fail("parsed"))
    assert_error(data, PDF, "pdf_page_limit", max_pdf_pages=1)


def test_pdf_rejects_deep_structure():
    def add_deep(writer):
        node = DictionaryObject()
        writer._root_object[NameObject("/Custom")] = node
        for _ in range(70):
            child = DictionaryObject()
            node[NameObject("/Child")] = child
            node = child

    assert_error(pdf_fixture(None, change=add_deep), PDF, "structure_limit_exceeded")


def test_pdf_stream_expansion_is_bounded_before_extraction(monkeypatch):
    def add_stream(writer):
        stream = DecodedStreamObject()
        stream.set_data(b" " * (extraction._PDF_STREAM_BYTES + 1))
        writer.pages[0][NameObject("/Contents")] = writer._add_object(stream.flate_encode())

    data = pdf_fixture(None, change=add_stream)
    monkeypatch.setattr(PageObject, "extract_text", lambda *args, **kwargs: pytest.fail("parsed"))
    assert_error(data, PDF, "structure_limit_exceeded")


def test_pdf_combined_content_stream_size_is_checked(monkeypatch):
    def add_streams(writer):
        references = []
        for _ in range(2):
            stream = DecodedStreamObject()
            stream.set_data(b" " * (extraction._PDF_STREAM_BYTES // 2 + 1))
            references.append(writer._add_object(stream.flate_encode()))
        writer.pages[0][NameObject("/Contents")] = ArrayObject(references)

    data = pdf_fixture(None, change=add_streams)
    monkeypatch.setattr(PageObject, "extract_text", lambda *args, **kwargs: pytest.fail("parsed"))
    assert_error(data, PDF, "structure_limit_exceeded")


def test_pdf_parser_configuration_is_restored():
    before = extraction.pypdf.get_configuration()
    extract_document(pdf_fixture("Text"), PDF, LIMITS)
    assert extraction.pypdf.get_configuration() == before


def test_docx_preserves_interleaved_paragraph_table_order_and_source_refs():
    result = extract_document(docx_fixture(), DOCX, LIMITS)
    assert [section.text for section in result.sections] == [
        "Plan",
        "First paragraph",
        "First table",
        "Second paragraph",
        "Second table\tBeside it",
        "Last row",
    ]
    assert [section.heading for section in result.sections] == [
        "Paragraph 1 · Plan",
        "Paragraph 2 · Plan",
        "Table 1, row 1 · Plan",
        "Paragraph 3 · Plan",
        "Table 2, row 1 · Plan",
        "Table 2, row 2 · Plan",
    ]
    assert all(
        section.page_number is None and section.line_start is None for section in result.sections
    )


def test_docx_nested_tables_are_retained_in_cell_source_order():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "Before"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    cell.add_paragraph("After")
    output = BytesIO()
    document.save(output)
    result = extract_document(output.getvalue(), DOCX, LIMITS)
    assert result.sections[0].text.split() == ["Before", "Nested", "After"]


def test_docx_heading_length_is_capped_without_losing_paragraph_reference():
    document = Document()
    document.add_heading("h" * 800, 1)
    document.add_paragraph("Body")
    output = BytesIO()
    document.save(output)
    result = extract_document(output.getvalue(), DOCX, LIMITS)
    assert len(result.sections[1].heading) == 500
    assert result.sections[1].heading.startswith("Paragraph 2 · ")


@pytest.mark.parametrize(
    "name",
    ["../outside", "/absolute", "word/../escape", "word\\evil", "C:/outside", "word//duplicate"],
)
def test_docx_rejects_unsafe_member_names(name):
    data = change_zip(docx_fixture(), {}, extras=[(name, b"bad")])
    if "\\" in name:
        # Windows ZipInfo normalizes separators when writing. Restore the
        # hostile filename in both archive headers for this fixture.
        data = data.replace(name.replace("\\", "/").encode(), name.encode())
    assert_error(data, DOCX, "unsafe_docx")


def test_docx_rejects_duplicate_member_names():
    with pytest.warns(UserWarning, match="Duplicate name"):
        data = change_zip(docx_fixture(), {}, extras=[("word/document.xml", b"bad")])
    assert_error(data, DOCX, "unsafe_docx")


@pytest.mark.parametrize(
    "name", ["word/vbaProject.bin", "word/embeddings/object.bin", "word/activeX/control.bin"]
)
def test_docx_rejects_macro_embedded_and_activex_parts(name):
    assert_error(change_zip(docx_fixture(), {}, extras=[(name, b"bad")]), DOCX, "unsafe_docx")


def test_docx_rejects_symbolic_link_members():
    entry = ZipInfo("word/link")
    entry.external_attr = 0o120777 << 16
    data = change_zip(docx_fixture(), {}, extras=[(entry, b"target")])
    assert_error(data, DOCX, "unsafe_docx")


def test_docx_rejects_ole_compound_file_with_misleading_part_name():
    data = change_zip(
        docx_fixture(),
        {},
        extras=[("word/arbitrary.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"contents")],
    )
    assert_error(data, DOCX, "unsafe_docx")


def test_docx_rejects_external_relationships_without_fetching():
    name = "word/_rels/document.xml.rels"
    with ZipFile(BytesIO(docx_fixture())) as source:
        rels = source.read(name).replace(
            b"</Relationships>",
            b'<Relationship Id="external" Type="http://example.test/hyperlink" '
            b'Target="https://never-fetch.invalid/secret" TargetMode="External"/></Relationships>',
        )
    assert_error(change_zip(docx_fixture(), {name: rels}), DOCX, "unsafe_docx")


def test_docx_rejects_internal_relationship_escaping_package():
    name = "_rels/.rels"
    with ZipFile(BytesIO(docx_fixture())) as source:
        rels = source.read(name).replace(b'Target="word/document.xml"', b'Target="../outside.xml"')
    assert_error(change_zip(docx_fixture(), {name: rels}), DOCX, "unsafe_docx")


@pytest.mark.parametrize(
    "payload",
    [
        b'<!DOCTYPE w:document [<!ENTITY secret SYSTEM "file:///never-read">]>',
        b'<!DOCTYPE w:document [<!ENTITY a "1234567890"><!ENTITY b "&a;&a;&a;&a;">]>',
    ],
)
def test_docx_rejects_dtd_and_entities(payload):
    with ZipFile(BytesIO(docx_fixture())) as source:
        main = source.read("word/document.xml")
    offset = main.index(b"?>") + 2
    main = main[:offset] + payload + main[offset:]
    assert_error(change_zip(docx_fixture(), {"word/document.xml": main}), DOCX, "unsafe_docx")


def test_docx_rejects_altchunk():
    with ZipFile(BytesIO(docx_fixture())) as source:
        main = source.read("word/document.xml").replace(b"<w:body>", b"<w:body><w:altChunk/>")
    assert_error(change_zip(docx_fixture(), {"word/document.xml": main}), DOCX, "unsafe_docx")


def test_docx_xml_checks_also_apply_to_misleading_file_extensions():
    with ZipFile(BytesIO(docx_fixture())) as source:
        types = source.read("[Content_Types].xml").replace(
            b"</Types>",
            b'<Override PartName="/word/unusual.bin" '
            b'ContentType="application/custom+xml"/></Types>',
        )
    data = change_zip(
        docx_fixture(),
        {"[Content_Types].xml": types},
        extras=[("word/unusual.bin", b'<!DOCTYPE x [<!ENTITY a "x">]><x/>')],
    )
    assert_error(data, DOCX, "unsafe_docx")


def test_docx_xml_content_type_overrides_match_members_case_insensitively():
    data = docx_fixture()
    with ZipFile(BytesIO(data)) as source:
        types = source.read("[Content_Types].xml").replace(b"/word/styles.xml", b"/word/style.bin")
        relationships = source.read("word/_rels/document.xml.rels").replace(
            b'Target="styles.xml"', b'Target="STYLE.BIN"'
        )
        styles = source.read("word/styles.xml").replace(
            b"?>", b'?>\n<!DOCTYPE w:styles [<!ENTITY bypass "forbidden-doctype">]>', 1
        )
    data = change_zip(
        data,
        {
            "[Content_Types].xml": types,
            "word/_rels/document.xml.rels": relationships,
            "word/styles.xml": None,
        },
        extras=[("word/STYLE.BIN", styles)],
    )
    assert_error(data, DOCX, "unsafe_docx")


@pytest.mark.parametrize("name", ["[Content_Types].xml", "_rels/.rels", "word/document.xml"])
def test_docx_requires_document_package_parts(name):
    assert_error(change_zip(docx_fixture(), {name: None}), DOCX, "invalid_docx")


def test_docx_rejects_other_openxml_document_types():
    with ZipFile(BytesIO(docx_fixture())) as source:
        types = source.read("[Content_Types].xml").replace(
            b"wordprocessingml.document.main+xml", b"presentationml.presentation.main+xml"
        )
    assert_error(change_zip(docx_fixture(), {"[Content_Types].xml": types}), DOCX, "invalid_docx")


def test_docx_expansion_and_entry_count_are_bounded():
    data = docx_fixture()
    assert_error(data, DOCX, "docx_expansion_limit", max_docx_entries=2)
    assert_error(data, DOCX, "docx_expansion_limit", max_docx_expanded_bytes=1_000)
    assert_error(data, DOCX, "docx_expansion_limit", max_docx_entry_bytes=1_000)
    bomb = change_zip(data, {}, extras=[("word/padding.bin", b"0" * 100_000)])
    assert_error(bomb, DOCX, "docx_expansion_limit")


@pytest.mark.parametrize("data", [b"not a zip", b"PK\x03\x04broken"])
def test_docx_rejects_invalid_archives(data):
    assert_error(data, DOCX, "invalid_docx")


def test_plain_text_preserves_real_line_ranges_and_utf8():
    result = extract_document("\ufeffFirst\r\nSecond\r\n\r\nनमस्ते\n".encode(), "text/plain", LIMITS)
    assert [
        (section.text, section.line_start, section.line_end) for section in result.sections
    ] == [("First\nSecond", 1, 2), ("नमस्ते", 4, 4)]
    assert all(section.page_number is None for section in result.sections)


def test_plain_text_groups_have_bounded_line_count():
    result = extract_document(("line\n" * 161).encode(), "text/plain", LIMITS)
    assert [(section.line_start, section.line_end) for section in result.sections] == [
        (1, 80),
        (81, 160),
        (161, 161),
    ]


@pytest.mark.parametrize("content_type", ["text/plain", "text/markdown"])
def test_text_strict_utf8_controls_and_byte_limits(content_type):
    assert_error(b"bad\xff", content_type, "invalid_utf8")
    assert_error(b"bad\x00", content_type, "unsafe_text")
    assert_error("bad\u0085".encode(), content_type, "unsafe_text")
    assert_error("éé".encode(), content_type, "text_limit_exceeded", max_text_bytes=3)


def test_markdown_uses_token_source_ranges_headings_and_plain_link_labels():
    data = (
        b"# Plan\n\nA **bold** [link](https://never-fetch.invalid).\ncontinued\n"
        b"\n## Code\n\n```py\nx = 1\n```\n"
    )
    result = extract_document(data, "text/markdown", LIMITS)
    assert [
        (section.text, section.heading, section.line_start, section.line_end)
        for section in result.sections
    ] == [
        ("Plan", "Plan", 1, 1),
        ("A bold link.\ncontinued", "Plan", 3, 4),
        ("Code", "Code", 6, 6),
        ("x = 1", "Code", 8, 10),
    ]


def test_markdown_disables_html_and_retains_image_alt_without_fetching():
    result = extract_document(
        b"<script>alert(1)</script>\n\n![A useful **image**](https://never-fetch.invalid/img.png)",
        "text/markdown",
        LIMITS,
    )
    assert [section.text for section in result.sections] == [
        "<script>alert(1)</script>",
        "A useful image",
    ]


def test_markdown_rejects_excessive_block_nesting():
    assert_error(b">" * 40 + b" nested", "text/markdown", "structure_limit_exceeded")


def test_output_section_and_pdf_text_limits_are_enforced():
    assert_error(b"a\n\nb", "text/plain", "section_limit_exceeded", max_sections=1)
    assert_error(pdf_fixture("Long text"), PDF, "text_limit_exceeded", max_text_bytes=3)


@pytest.mark.parametrize("content_type", ["text/plain", "text/markdown"])
def test_whitespace_only_documents_have_explicit_empty_warning(content_type):
    result = extract_document(b" \n\t\n", content_type, LIMITS)
    assert result.sections == []
    assert result.warning_code == "no_extractable_text"


def test_timeout_has_safe_terminal_code(monkeypatch):
    calls = iter([0, 100, 100, 100])
    monkeypatch.setattr(extraction.time, "monotonic", lambda: next(calls, 100))
    error = assert_error(b"text", "text/plain", "extraction_timeout")
    assert not error.retryable


def test_unknown_content_type_has_stable_error_code():
    assert_error(b"text", "application/octet-stream", "unsupported_type")


@pytest.mark.parametrize(
    "values",
    [
        {"text": " "},
        {"text": "nul\x00"},
        {"text": "a", "page_number": 0},
        {"text": "a", "line_start": 2},
        {"text": "a", "line_start": 3, "line_end": 2},
        {"text": "a", "heading": " "},
        {"text": "a", "line_start": 2**31, "line_end": 2**31},
    ],
)
def test_section_model_rejects_invalid_publication_values(values):
    with pytest.raises(ValidationError):
        Section(**values)


@pytest.mark.parametrize(
    "values",
    [
        {"warning_code": "bad\n"},
        {"parser_version": " "},
        {"parser_version": "bad\x00"},
    ],
)
def test_result_model_rejects_invalid_publication_codes(values):
    with pytest.raises(ValidationError):
        ExtractionResult(sections=[], **values)
