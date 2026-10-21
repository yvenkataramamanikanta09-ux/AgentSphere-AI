"""Run real browser tests in a disposable local PostgreSQL database.

Uses the existing ignored local cluster configuration, never the development DB.
Only the database created by this run is dropped. Ignored HTML reports may include
generated disposable credentials in browser action labels.
"""

import json
import os
import re
import secrets
import shutil
import socket
import subprocess
import sys
import sysconfig
import time
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import UUID, uuid4
from zipfile import ZIP_STORED, ZipFile

import httpx
import psycopg
from docx import Document as WordDocument
from psycopg import sql
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select, update
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session

# isort: split
from agentsphere.config import BACKEND_DIR, Settings
from agentsphere.database import build_engine
from agentsphere.document_lifecycle import cleanup_job_id
from agentsphere.models import (
    AuthSession,
    Document,
    DocumentFile,
    DocumentJob,
    DocumentSection,
    JobKind,
    Membership,
    User,
    Workspace,
)
from agentsphere.security import hash_password
from agentsphere.storage import validate_storage_key

FRONTEND = Path(__file__).resolve().parents[1]
ORIGIN = "http://127.0.0.1:5174"


def python_command(*, module=None, script=None, arguments=(), purelib=None):
    """Use a real Python process on Windows, not a venv redirector parent."""
    if os.name == "nt":
        target = module or str(script)
        operation = (
            f"runpy.run_module({module!r},run_name='__main__')"
            if module
            else f"runpy.run_path({str(script)!r},run_name='__main__')"
        )
        bootstrap = (
            "import runpy,site,sys;"
            f"site.addsitedir({(purelib or sysconfig.get_paths()['purelib'])!r});"
            f"sys.argv={[target, *arguments]!r};{operation}"
        )
        return [sys._base_executable, "-I", "-c", bootstrap]
    return [sys.executable, *(["-m", module] if module else [str(script)]), *arguments]


def isolated_environment():
    """Guard every test helper before opening a database or touching private files."""
    url = make_url(os.environ["DATABASE_URL"])
    require_local_url(url)
    if os.environ.get("E2E_ISOLATED") != "1" or not re.fullmatch(
        r"agentsphere_e2e_[0-9a-f]{32}", url.database or ""
    ):
        raise SystemExit("Test helper requires this runner's disposable E2E database.")
    root = Path(os.environ.get("E2E_RUNTIME_ROOT", ""))
    private = Path(os.environ.get("PRIVATE_STORAGE_ROOT", ""))
    if not root.is_absolute() or not private.is_absolute():
        raise SystemExit("Test helper requires an absolute isolated private storage root.")
    root, private = root.resolve(), private.resolve()
    if private != root / "private" or not root.name.startswith("agentsphere_e2e_"):
        raise SystemExit("Test helper refuses storage outside this run's temporary directory.")
    marker = json.loads((root / "e2e-run.json").read_text(encoding="utf-8"))
    if marker != {"database": url.database, "private_root": str(private)}:
        raise SystemExit("Test helper refuses a mismatched E2E run marker.")
    return url, root, private


def atomic_json(path, value):
    temporary = path.with_name(path.name + "." + uuid4().hex + ".tmp")
    temporary.write_text(json.dumps(value), encoding="utf-8")
    deadline = time.monotonic() + 2
    try:
        while True:
            try:
                temporary.replace(path)
                return
            except PermissionError as error:
                # Windows readers may briefly hold a handle without delete
                # sharing. Preserve atomic publication and retry only that
                # bounded sharing/access race in this disposable mailbox.
                if (
                    os.name != "nt"
                    or error.winerror not in {5, 32, 33}
                    or time.monotonic() >= deadline
                ):
                    raise
                time.sleep(0.02)
    finally:
        temporary.unlink(missing_ok=True)


def worker_control(action, environment=None):
    if environment is not None:
        # Keep diagnostics private but available when the helper cannot start.
        with (FRONTEND / ".local/e2e/worker-control.log").open("a", encoding="utf-8") as log:
            return subprocess.run(
                python_command(script=Path(__file__), arguments=("--worker-control", action)),
                cwd=FRONTEND,
                env=environment,
                check=True,
                stdout=log,
                stderr=log,
                timeout=20,
            )
    _, root, _ = isolated_environment()
    atomic_json(root / "worker-request.json", {"action": action})
    expected = "paused" if action == "pause" else "running"
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        state_file = root / "worker-state.json"
        if state_file.is_file():
            state = json.loads(state_file.read_text(encoding="utf-8"))
            if state["state"] == expected:
                return 0
        time.sleep(0.1)
    raise SystemExit("Isolated worker did not acknowledge the requested state.")


def worker_supervisor():
    _, root, _ = isolated_environment()
    process = None
    state = None
    try:
        while True:
            action = json.loads((root / "worker-request.json").read_text())["action"]
            if action == "pause" and process is not None:
                process.terminate()
                try:
                    process.wait(timeout=8)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait()
                process = None
            if action == "resume" and process is None:
                process = subprocess.Popen(
                    python_command(
                        module="agentsphere.worker", purelib=os.environ["E2E_VENV_PURELIB"]
                    ),
                    cwd=BACKEND_DIR,
                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
                )
            if process is not None and process.poll() is not None:
                raise RuntimeError("Isolated processing worker exited unexpectedly.")
            current = "running" if process is not None else "paused"
            if current != state:
                atomic_json(root / "worker-state.json", {"state": current})
                state = current
            time.sleep(0.1)
    finally:
        if process is not None:
            process.terminate()
            process.wait(timeout=10)


def inspect_upload(document_id):
    url, _, private = isolated_environment()
    identifier = UUID(document_id)
    engine = build_engine(url.render_as_string(hide_password=False))
    try:
        with Session(engine) as db:
            attachment = db.scalar(
                select(DocumentFile).where(DocumentFile.document_id == identifier)
            )
            cleanup = db.get(DocumentJob, cleanup_job_id(identifier))
            key = attachment.storage_key if attachment else cleanup.storage_key if cleanup else None
            exists = (private / validate_storage_key(key)).is_file() if key else False
            print(
                json.dumps(
                    {
                        "document_count": db.scalar(
                            select(func.count())
                            .select_from(Document)
                            .where(Document.id == identifier)
                        ),
                        "file_count": int(attachment is not None),
                        "section_count": db.scalar(
                            select(func.count())
                            .select_from(DocumentSection)
                            .where(DocumentSection.document_id == identifier)
                        ),
                        "extract_job_count": db.scalar(
                            select(func.count())
                            .select_from(DocumentJob)
                            .where(
                                DocumentJob.document_id == identifier,
                                DocumentJob.kind == JobKind.EXTRACT,
                            )
                        ),
                        "private_original_exists": exists,
                        "cleanup_state": cleanup.state.value if cleanup else None,
                    }
                )
            )
    finally:
        engine.dispose()


def create_upload_fixtures(root):
    """Generate harmless documents at runtime; never copy personal documents."""
    folder = root / "synthetic"
    folder.mkdir()
    payloads = {
        "txt": (
            "synthetic-source.txt",
            b"\n\n".join(
                f"Synthetic source paragraph {number:02}.".encode() for number in range(1, 66)
            )
            + b"\n",
        ),
        "md": (
            "synthetic-source.md",
            (
                b"# Synthetic Markdown evidence\n\nA generated Markdown paragraph.\n\n"
                b"<script>window.uploadPreviewExecuted=true</script>\n\n"
                b"![Generated image label](https://example.invalid/never-fetch.png)\n"
            ),
        ),
        "malformed": (
            "synthetic-invalid.pdf",
            b"%PDF-1.7\nGenerated malformed PDF.\n%%EOF\n",
        ),
        "invalid_txt": ("synthetic-invalid-utf8.txt", b"Generated invalid UTF-8: \xff\n"),
        "invalid_md": ("synthetic-invalid-utf8.md", b"# Generated invalid UTF-8\n\xff\n"),
    }
    writer = PdfWriter()
    page = writer.add_blank_page(width=250, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 24 100 Td (Synthetic PDF evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    payloads["pdf"] = ("synthetic-source.pdf", output.getvalue())
    word = WordDocument()
    word.add_heading("Synthetic Word evidence", 1)
    word.add_paragraph("A generated DOCX paragraph.")
    row = word.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "Source", "Generated evidence"
    output = BytesIO()
    word.save(output)
    payloads["docx"] = ("synthetic-source.docx", output.getvalue())
    # Basic ZIP evidence passes intake; malformed XML must fail in the worker,
    # leaving previews and original downloads unavailable.
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_STORED) as package:
        package.writestr("[Content_Types].xml", b"<Types>")
        package.writestr("word/document.xml", b"<document>")
    payloads["malformed_docx"] = ("synthetic-invalid.docx", output.getvalue())
    result = {}
    for kind, (filename, payload) in payloads.items():
        path = folder / filename
        path.write_bytes(payload)
        result[kind] = str(path)
    return result


def require_local_url(url):
    # SQLAlchemy applies query parameters after URL fields; dbname/host overrides
    # would otherwise bypass the disposable database and loopback checks.
    if url.host not in {"localhost", "127.0.0.1", "::1"} or url.query:
        raise SystemExit("E2E database URLs must be loopback URLs without query parameters.")


def wait_ready(process, url):
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError("E2E server exited; inspect ignored .local/e2e logs.")
        try:
            if httpx.get(url, timeout=1).status_code == 200:
                return
        except httpx.TransportError:
            pass
        time.sleep(0.2)
    raise RuntimeError("E2E server did not become ready.")


def expire_user(user_id):
    """Test helper: only a marked account in this runner's disposable database."""
    url, _, _ = isolated_environment()
    engine = build_engine(url.render_as_string(hide_password=False))
    try:
        with Session(engine) as db, db.begin():
            user = db.get(User, UUID(user_id))
            if user is None or not user.email.startswith("stage4-expiry-"):
                raise SystemExit("Expiry helper refuses unmarked users.")
            db.execute(
                update(AuthSession)
                .where(AuthSession.user_id == user.id)
                .values(expires_at=datetime.now(UTC) - timedelta(seconds=1))
            )
    finally:
        engine.dispose()


def seed(url, password):
    engine = build_engine(url)
    result = {"users": {}}
    try:
        with Session(engine) as db, db.begin():
            users = {}
            roles = ("owner", "editor", "viewer", "expiry") + tuple(
                f"layout-{width}" for width in (1440, 1024, 768, 390, 320)
            )
            for role in roles:
                user = User(
                    email=f"stage4-{role}-{uuid4().hex}@example.com",
                    display_name=f"Browser {role.title()}",
                    password_hash=hash_password(password),
                )
                db.add(user)
                db.flush()
                users[role] = user
                personal = Workspace(name="Personal workspace", owner_user_id=user.id)
                db.add(personal)
                db.flush()
                result["users"][role] = {
                    "id": str(user.id),
                    "email": user.email,
                    "display_name": user.display_name,
                }
            workspace = Workspace(name="Browser research", owner_user_id=users["owner"].id)
            db.add(workspace)
            db.flush()
            result["workspace_id"] = str(workspace.id)
            for role in ("editor", "viewer", *roles[4:]):
                db.add(
                    Membership(
                        workspace_id=workspace.id,
                        user_id=users[role].id,
                        role="viewer" if role == "viewer" else "editor",
                    )
                )
            for index in range(12):
                db.add(
                    Document(
                        workspace_id=workspace.id,
                        filename=f"seed-{index:02}.pdf",
                        content_type="application/pdf",
                        size_bytes=2048,
                        status="archived" if index == 0 else "registered",
                        created_at=datetime.now(UTC) + timedelta(seconds=index),
                    )
                )
        return result
    finally:
        engine.dispose()


def main():
    if len(sys.argv) == 3 and sys.argv[1] == "--expire-user":
        expire_user(sys.argv[2])
        return 0
    if len(sys.argv) == 3 and sys.argv[1] == "--inspect-upload":
        inspect_upload(sys.argv[2])
        return 0
    if len(sys.argv) == 3 and sys.argv[1] == "--worker-control":
        if sys.argv[2] not in {"pause", "resume"}:
            raise SystemExit("Worker control accepts only pause or resume.")
        return worker_control(sys.argv[2])
    if sys.argv[1:] == ["--worker-supervisor"]:
        return worker_supervisor()
    settings = Settings()
    if settings.environment != "development" or settings.test_database_url is None:
        raise SystemExit("Configure the existing local development TEST_DATABASE_URL first.")
    test_url = make_url(settings.test_database_url.get_secret_value())
    require_local_url(test_url)
    if test_url.host not in {"localhost", "127.0.0.1", "::1"} or not (
        test_url.database or ""
    ).endswith("_test"):
        raise SystemExit("Only a loopback database ending in _test is supported.")
    cluster = json.loads((BACKEND_DIR / ".local/postgres-config.json").read_text())
    if cluster["port"] != test_url.port:
        raise SystemExit("Local cluster and test database ports differ; nothing changed.")
    for port in (8002, 5174):
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", port))
    node = shutil.which("node")
    cli = FRONTEND / "node_modules/@playwright/test/cli.js"
    if not node or not cli.is_file():
        raise SystemExit("Run npm ci and npx playwright install chromium first.")
    name = "agentsphere_e2e_" + uuid4().hex
    url = test_url.set(database=name).render_as_string(hide_password=False)
    temporary = TemporaryDirectory(prefix="agentsphere_e2e_")
    runtime_root = Path(temporary.name).resolve()
    private_root = runtime_root / "private"
    atomic_json(
        runtime_root / "e2e-run.json",
        {"database": name, "private_root": str(private_root)},
    )
    env = os.environ.copy()
    installed_chrome = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")
    if "E2E_BROWSER_CHANNEL" not in env and installed_chrome.is_file():
        env["E2E_BROWSER_CHANNEL"] = "chrome"
    env.update(
        DATABASE_URL=url,
        APP_ORIGIN=ORIGIN,
        ALLOW_INSECURE_LOCAL_HTTP="true",
        ENVIRONMENT="development",
        MAX_UPLOAD_BYTES="10000000",
        AUTH_RATE_LIMIT_KEY=secrets.token_urlsafe(48),
        E2E_PASSWORD=secrets.token_urlsafe(30),
        E2E_PYTHON=sys.executable,
        E2E_VENV_PURELIB=sysconfig.get_paths()["purelib"],
        E2E_ISOLATED="1",
        E2E_RUNTIME_ROOT=str(runtime_root),
        E2E_UPLOAD_FIXTURES=json.dumps(create_upload_fixtures(runtime_root)),
        PRIVATE_STORAGE_ROOT=str(private_root),
        WORKER_POLL_SECONDS="0.1",
        LOGIN_ACCOUNT_LIMIT="100",
        LOGIN_IP_LIMIT="1000",
        CSRF_IP_LIMIT="1000",
        BACKEND_TARGET="http://127.0.0.1:8002",
    )
    admin = {
        "host": test_url.host,
        "port": test_url.port,
        "user": "agentsphere_admin",
        "password": cluster["admin_password"],
        "dbname": "postgres",
        "autocommit": True,
        "connect_timeout": 3,
    }
    created = False
    servers = []
    logs = []
    supervisor = None
    try:
        with psycopg.connect(**admin) as db:
            db.execute(
                sql.SQL("CREATE DATABASE {} OWNER {} TEMPLATE template0").format(
                    sql.Identifier(name), sql.Identifier(test_url.username)
                )
            )
            created = True
        print("Created a unique disposable E2E database; development records are untouched.")
        subprocess.run(
            python_command(module="alembic", arguments=("upgrade", "head")),
            cwd=BACKEND_DIR,
            env=env,
            check=True,
        )
        subprocess.run(
            python_command(module="alembic", arguments=("check",)),
            cwd=BACKEND_DIR,
            env=env,
            check=True,
        )
        env["E2E_FIXTURE"] = json.dumps(seed(url, env["E2E_PASSWORD"]))
        log_dir = FRONTEND / ".local/e2e"
        log_dir.mkdir(parents=True, exist_ok=True)
        log = (log_dir / "backend.log").open("w", encoding="utf-8")
        logs.append(log)
        server = subprocess.Popen(
            python_command(
                module="uvicorn",
                arguments=(
                    "agentsphere.main:create_app",
                    "--factory",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    "8002",
                    "--no-proxy-headers",
                ),
            ),
            cwd=BACKEND_DIR,
            env=env,
            stdout=log,
            stderr=log,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        servers.append(server)
        wait_ready(server, "http://127.0.0.1:8002/health/live")
        frontend_log = (log_dir / "frontend.log").open("w", encoding="utf-8")
        logs.append(frontend_log)
        vite = subprocess.Popen(
            [
                node,
                str(FRONTEND / "node_modules/vite/bin/vite.js"),
                "--host",
                "127.0.0.1",
                "--port",
                "5174",
                "--strictPort",
            ],
            cwd=FRONTEND,
            env=env,
            stdout=frontend_log,
            stderr=frontend_log,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        servers.append(vite)
        wait_ready(vite, ORIGIN)
        subprocess.run(
            [sys.executable, "scripts/verify_document_api.py"],
            cwd=FRONTEND,
            env=env | {"VERIFICATION_ORIGIN": ORIGIN},
            check=True,
        )
        # Keep the legacy API verification independent of background work. Only
        # after it passes, give browser uploads a real separate Python worker.
        worker_log = (log_dir / "worker.log").open("w", encoding="utf-8")
        logs.append(worker_log)
        atomic_json(runtime_root / "worker-request.json", {"action": "resume"})
        supervisor = subprocess.Popen(
            python_command(script=Path(__file__), arguments=("--worker-supervisor",)),
            cwd=FRONTEND,
            env=env,
            stdout=worker_log,
            stderr=worker_log,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        servers.append(supervisor)
        worker_control("resume", env)
        return subprocess.run(
            [node, str(cli), "test", *sys.argv[1:]], cwd=FRONTEND, env=env, check=False
        ).returncode
    finally:
        worker_paused = False
        if supervisor is not None and supervisor.poll() is None:
            try:
                worker_control("pause", env)
                worker_paused = True
            except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
                print("Isolated worker shutdown required forced process cleanup.")
        if (
            os.name == "nt"
            and supervisor is not None
            and supervisor.poll() is None
            and not worker_paused
        ):
            subprocess.run(
                ["taskkill", "/PID", str(supervisor.pid), "/T", "/F"],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                creationflags=subprocess.CREATE_NO_WINDOW,
            )
        for server in reversed(servers):
            if server.poll() is not None:
                continue
            server.terminate()
            try:
                server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                server.kill()
                server.wait()
        for log in logs:
            log.close()
        try:
            if created:
                # name is generated above, never taken from user input or a saved fixture.
                with psycopg.connect(**admin) as db:
                    db.execute(
                        sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(name))
                    )
                print("Removed only this run's disposable database.")
        finally:
            temporary.cleanup()
            print("Removed this run's synthetic fixtures and isolated private storage.")


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (psycopg.Error, subprocess.CalledProcessError):
        # Driver exception text can contain connection details; keep it out of reports.
        raise SystemExit(
            "Local E2E setup failed; check PostgreSQL and the documented prerequisites."
        ) from None
